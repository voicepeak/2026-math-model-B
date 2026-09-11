from pathlib import Path
import sys,json,hashlib,re
ROOT=Path(__file__).resolve().parents[3];W=ROOT/'B题/work/paper'
sys.path.insert(0,'/Users/a1/.codex/skills/math-modeling/tools/docx/scripts')
import paper_format as pf
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
pdf=W/'render_checked/完整论文.pdf';doc=Document(ROOT/'完整论文.docx');r=PdfReader(pdf)
texts=[p.extract_text() for p in r.pages]
appendix=next(i+1 for i,t in enumerate(texts) if re.match(r'^\s*\d+\s*附录\s',t))
issues=pf.validate_paper_structure(doc,'cumcm',min_content_units=7000,min_equations=20,min_figures=6,min_tables=7,rendered_pages=len(r.pages),target_pages=12,official_max_pages=30)
assert '关键词' in texts[0] and '1 问题重述' not in texts[0], 'Abstract must occupy first page only'
assert any('1 问题重述' in t for t in texts[1:3])
assert appendix-2<=30
assert all(abs(float(p.mediabox.width)-595.3)<1 and abs(float(p.mediabox.height)-841.9)<1 for p in r.pages)
assert all(len(t.strip())>15 for t in texts),'Blank page'
raw='\n'.join(p.text for p in doc.paragraphs)
for s in ['[待补充]','Subagent','门禁','Checkpoint','202619033012','何林军']:
 assert s not in raw,s
for pre,nxt in [('5 问题三','6 问题四'),('6 问题四','7 模型评价')]:
 a=next(i for i,p in enumerate(doc.paragraphs) if p.text.startswith(pre));b=next(i for i,p in enumerate(doc.paragraphs[a+1:],a+1) if p.text.startswith(nxt))
 assert all(not p.text.strip() or p.style.name.startswith('Heading') for p in doc.paragraphs[a+1:b])
for path in [ROOT/'完整论文.docx',pdf,ROOT/'AI工具使用详情.pdf']:assert path.stat().st_size<20*1024**2
snapshot=json.loads((W/'input_snapshot.json').read_text())
assert all(hashlib.sha256((ROOT/f).read_bytes()).hexdigest()==sha for f,sha in snapshot.items()),'Upstream changed'
fig_paths=[W/'figures/raw_q1_observations.png',ROOT/'figures/result_q1_counter.png',ROOT/'figures/process_q1_constraints.png',W/'figures/raw_q2_prior.png',ROOT/'figures/process_q2_score.png',ROOT/'figures/result_q2_candidates.png']
from PIL import Image
assert all(Image.open(p).width/(13.2/2.54)>=300 for p in fig_paths)
fonts={}
for page in r.pages:
 for f in page['/Resources'].get('/Font',{}).get_object().values():
  f=f.get_object();name=str(f.get('/BaseFont'));fs=f.get('/DescendantFonts',[f])
  for ff in fs:
   fd=ff.get_object().get('/FontDescriptor')
   if fd:fonts[name]=any(x in fd.get_object() for x in ['/FontFile','/FontFile2','/FontFile3'])
assert all(fonts.values()),fonts
report={'scope':'问题一、二论文草稿；三、四及完整源码附录保留框架，非全题提交版','override_source':'用户2026-09-11明确要求先写1/2，3/4留空','thresholds':{'min_content_units':7000,'min_equations':20,'min_figures':6,'min_tables':7,'target_pages':12},'metrics':{'content_units':pf._content_units('\n'.join(pf._document_texts(doc))),'total_pages':len(r.pages),'abstract_pages':1,'body_pages':appendix-2,'appendix_start_page':appendix,'equations':len(doc._element.findall('.//'+qn('m:oMath'))),'figures':len(doc.inline_shapes),'tables':len(doc.tables),'min_effective_dpi':min(Image.open(p).width/(13.2/2.54) for p in fig_paths)},'issues':issues,'fonts_embedded':fonts,'source_snapshot_unchanged':True,'hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [ROOT/'完整论文.docx',pdf,ROOT/'AI工具使用详情.pdf',W/'build_paper.py']},'passed':not issues}
(W/'quality_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2));sys.exit(bool(issues))
