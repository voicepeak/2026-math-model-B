# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"C:\Users\张靖浩\Desktop\2026数学建模\2026-math-model-B\B题\work\paper\AI工具使用详情_20260913.docx"
doc = Document()

def add(text, size=12, bold=False, center=False, indent=True):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.line_spacing = Pt(18)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.space_after = Pt(0)
    if indent and not center:
        pf.first_line_indent = Pt(24)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    rPr = r._r.get_or_add_rPr()
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")
    rf.set(qn("w:eastAsia"), "宋体")
    rPr.insert(0, rf)
    return p

add("AI工具使用详情", size=16, bold=True, center=True)
add("一、工具、版本及记录范围", bold=True, indent=False)
add("本项目竞赛期间的AI辅助通过多轮会话完成。已记录的工具与系统标识包括：OpenAI Codex（会话系统标识GPT-6，"
    "2026年9月11日、12日两轮），以及OpenCode命令行工具（模型标识deepseek-v4-flash-vision-exp，2026年9月13日一轮）。"
    "记录日期截至2026年9月13日。此前项目历史中以Gemini等命名的算法版本无法仅凭文件夹名称确认具体工具型号，"
    "应由参赛队按实际记录核实补充。")
add("二、各轮使用目的与环节", bold=True, indent=False)
add("1. 2026年9月11日：依据已通过检查的问题一、二成果整理推导、撰写论文草稿、统一符号与数值、"
    "构建Word原生公式及表格，并检查导出后的版面；图表副本仅调整了轴向坐标色条标注和检测点形状标记。")
add("2. 2026年9月12日：依据最终ultimate200代码、真实结果表与已有图件续写第三题，补充覆盖证明、"
    "算法说明、统计口径与适用边界，更新摘要和结论，并构建公式、可编辑表格及导出PDF。")
add("3. 2026年9月13日：按评委视角对全稿进行结构补全与表达优化。具体包括：(1) 依据现有问题四代码"
    "（B题/code/strategy.py）与官方演练日志（q4_五次演练汇总/）补写第6章问题四全文，含定向可见性模型、"
    "600 m网格四角覆盖引理、认证清除与条带回退、演练结果分析，以及7.4与8.3；(2) 生成14幅新图件并重绘"
    "问题三图件为中文标注（生成脚本B题/work/paper_opt/make_figures.py），包括总体技术路线图、问题一区域状态图、"
    "问题二保证接收区与选点对照图、问题三覆盖证明与清除进度图、问题四可见性、演练汇总、动作轨迹与时间分解图；"
    "(3) 插入图件、统一图号与交叉引用、新增表14—表16、补充附录清单；(4) 更新摘要。")
add("三、主要提示与处理过程", bold=True, indent=False)
add("前两轮的提示分别为“先将问题一、二已通过的成果完整写入并搭好全文框架”和“第三题已完成，"
    "用ultimate200继续写论文，第四题还在跑”。本轮主要提示为：“从零开始用评委的角度审查目前的论文问题”，"
    "以及“把能做到都做了，不只是用闲置的图，而是什么好就用什么、生成什么来用，请优化”。"
    "处理时通读全文与结果目录，清点图数与空白章节，核对问题四五场官方演练日志，重建动作耗时分解，"
    "按上述清单生成并插入内容。")
add("四、采纳、修改与核验情况", bold=True, indent=False)
add("新增内容中的全部数值取自现有结果与日志文件（study_summary.json、final_holdout.json、"
    "q4_五次演练汇总/*/summary.json与client.jsonl等），未编造官方成绩；问题三、问题四正式测试尚未执行，"
    "论文以“未执行”明示。机器辅助核验包括图件数据与结果文件的对应、图号引用完整性检查、"
    "DOCX结构与PDF渲染分页检查。第四题正文与新增图件为AI按现有实现整理生成，尚未经参赛队逐项人工审查；"
    "本说明不将机器核验冒称为人工核验。")
add("五、后续补充事项", bold=True, indent=False)
add("待执行问题三、问题四各三次正式测试后，应回填论文表13与表16，导出加密日志，"
    "并同步更新本说明的AI参与环节与结果核验记录；最终提交前由参赛队核实本说明与论文AI工具使用声明、"
    "实际使用过程一致。本文件不含姓名、队号、学校或赛区信息。")

doc.save(OUT)
print("saved", OUT)
